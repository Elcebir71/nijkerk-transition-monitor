import openpyxl, json, csv
from openpyxl.styles import Font, PatternFill, Alignment
from constants import NORMS, BIOGAS_YIELD, MEST_DENSITY_TON_PER_M3

ARIAL = "Arial"
TITLE = Font(name=ARIAL, bold=True, size=13)
SECTION_FONT = Font(name=ARIAL, bold=True, size=11, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="2E5B7A")
HEADER_FONT = Font(name=ARIAL, bold=True, size=10)
HEADER_FILL = PatternFill("solid", fgColor="D9E6EC")
BODY = Font(name=ARIAL, size=10)
NOTE_FONT = Font(name=ARIAL, italic=True, size=9, color="666666")
WARN_FONT = Font(name=ARIAL, bold=True, size=11, color="B00000")
WARN_FILL = PatternFill("solid", fgColor="FCE4E4")
WRAP = Alignment(wrap_text=True, vertical="top")

summary = json.load(open("../data/generation_summary.json"))

wb = openpyxl.Workbook()

# --- Sheet 1: Samenvatting ---
ws = wb.active
ws.title = "Samenvatting"
ws.column_dimensions["A"].width = 100
ws["A1"] = "Nijkerk Agricultural Transition Monitor - SYNTHETISCHE demodataset (75 fictieve bedrijven)"
ws["A1"].font = TITLE
r = 2
warn = ws.cell(row=r, column=1, value="LET OP: alle 75 bedrijven, namen, locaties en individuele dieraantallen in dit bestand zijn VOLLEDIG "
                                       "FICTIEF. Geen enkele rij komt overeen met een echt adres, echt BAG-perceel, echt BRP-perceel of bestaand "
                                       "bedrijf. (v1.1, 29-09-2026) De locaties zijn VOLLEDIG willekeurig gegenereerd binnen een schematische "
                                       "Nijkerk-omtrek (buiten het bebouwde centrum, min. 150m onderlinge afstand) - GEEN locatie is afgeleid van "
                                       "een echt perceel of gebouw. Uitsluitend de TOTALEN per diercategorie zijn gekalibreerd op (een "
                                       "veiligheidsmarge onder) de echte CBS-Landbouwtelling 2023 voor gemeente Nijkerk.")
warn.font = WARN_FONT
warn.fill = WARN_FILL
warn.alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=1)
ws.row_dimensions[r].height = 55
r += 2

ws.cell(row=r, column=1, value="Doel: eerste synthetische dataset voor een prototype 'Nijkerk Agricultural Transition Monitor' "
                                "(kaart + scenario-rekenmachine), zodat een werkend demo gebouwd en getoond kan worden zonder "
                                "echte, mogelijk herleidbare bedrijfsgegevens te publiceren. Zodra gevalideerde echte data "
                                "beschikbaar/vrijgegeven is, vervangt die dit bestand 1-op-1 (zelfde kolomstructuur).").font = BODY
ws.cell(row=r, column=1).alignment = WRAP
ws.row_dimensions[r].height = 45
r += 2

section = ws.cell(row=r, column=1, value="Totalen: synthetisch vs. doel (85% van echt 2023) vs. echt CBS 2023")
section.font = SECTION_FONT
section.fill = SECTION_FILL
r += 1
ws.cell(row=r, column=1, value="Categorie").font = HEADER_FONT
for c, lab in zip("BCD", ["Synthetisch (75 bedr.)", "Doel (85% van 2023)", "Echt CBS 2023 (GM0267)"]):
    ws[f"{c}{r}"] = lab
    ws[f"{c}{r}"].font = HEADER_FONT
for c in "ABCD":
    ws[f"{c}{r}"].fill = HEADER_FILL
r += 1
for cat, real_val in summary["real_2023"].items():
    ws.cell(row=r, column=1, value=cat)
    ws.cell(row=r, column=2, value=summary["check"][cat])
    ws.cell(row=r, column=3, value=round(real_val * summary["target_fraction"]))
    ws.cell(row=r, column=4, value=real_val)
    r += 1
r += 1
ws.cell(row=r, column=1, value=f"Totaal N-excretie (synthetisch): {summary['total_n_kg']:,.0f} kg N/jaar").font = Font(name=ARIAL, bold=True)
r += 1
ws.cell(row=r, column=1, value=f"Totaal mest (synthetisch): {summary['total_manure_ton']:,.0f} ton/jaar").font = Font(name=ARIAL, bold=True)
r += 1
ws.cell(row=r, column=1, value=f"Totaal geschat biogaspotentieel (synthetisch): {summary['total_biogas_m3']:,.0f} m3/jaar").font = Font(name=ARIAL, bold=True)
r += 2

note = ws.cell(row=r, column=1, value="Ter vergelijking: het hoofdrapport (Nijkerk_Emissiereductie_Rapport.pdf) gebruikte een specifieke, "
    "kleinere deelverzameling van 81 ECHTE bedrijven nabij Natura 2000/Arkemheen (RVO Woo-verzoek, 2022-dieraantallen, 173.295 ton mest, "
    "899.245 kg N). Deze synthetische dataset dekt de HELE gemeente Nijkerk (alle diersoorten, CBS 2023) en is dus qua opzet en omvang een "
    "ANDER, breder cijfer - niet rechtstreeks vergelijkbaar met de 81-bedrijven-analyse.")
note.font = NOTE_FONT
note.alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=1)
ws.row_dimensions[r].height = 60

# --- Sheet 2: Bronnen en normen ---
ws2 = wb.create_sheet("Bronnen en normen")
ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 14
ws2.column_dimensions["D"].width = 90
ws2["A1"] = "Per-diercategorie normen gebruikt voor de synthetische berekening"
ws2["A1"].font = TITLE
r = 3
for c, lab in zip("ABCD", ["Categorie", "kg N/dier/jr", "m3 mest/dier/jr", "Bron / toelichting"]):
    ws2[f"{c}{r}"] = lab
    ws2[f"{c}{r}"].font = HEADER_FONT
    ws2[f"{c}{r}"].fill = HEADER_FILL
r += 1
NORMS_TEXT = [(cat, n_kg, m3, src) for cat, (n_kg, m3, _mtype, src) in NORMS.items()]
for cat, nkg, m3, src in NORMS_TEXT:
    ws2.cell(row=r, column=1, value=cat).font = BODY
    ws2.cell(row=r, column=2, value=nkg).font = BODY
    ws2.cell(row=r, column=3, value=m3).font = BODY
    c = ws2.cell(row=r, column=4, value=src)
    c.font = NOTE_FONT
    c.alignment = WRAP
    ws2.row_dimensions[r].height = 28
    r += 1
r += 1
ws2.cell(row=r, column=1, value="Biogasopbrengst per mesttype (m3 biogas/ton) - typische literatuurwaarden, niet individueel geverifieerd:").font = Font(name=ARIAL, bold=True)
r += 1
for mtype, (yieldval, src) in BIOGAS_YIELD.items():
    ws2.cell(row=r, column=1, value=mtype).font = BODY
    ws2.cell(row=r, column=2, value=yieldval).font = BODY
    c = ws2.cell(row=r, column=4, value=src)
    c.font = NOTE_FONT
    c.alignment = WRAP
    r += 1
r += 1
ws2.cell(row=r, column=1, value="Mestdichtheid per mesttype (ton/m3) - APPROX, niet uit een specifieke RVO/CBS-tabel:").font = Font(name=ARIAL, bold=True)
r += 1
for mtype, dens in MEST_DENSITY_TON_PER_M3.items():
    ws2.cell(row=r, column=1, value=mtype).font = BODY
    ws2.cell(row=r, column=2, value=dens).font = BODY
    c = ws2.cell(row=r, column=4, value="Drijfmest (rundvee/varkens/overig) is overwegend water, ~1,0 ton/m3. Vaste pluimveemest heeft een "
                                          "hoger drogestofgehalte en een merkbaar lagere dichtheid (~0,65 ton/m3 hier gebruikt) - dus m3 en ton "
                                          "zijn voor pluimveebedrijven NIET gelijk, in tegenstelling tot een eerdere versie van dit bestand.")
    c.font = NOTE_FONT
    c.alignment = WRAP
    ws2.row_dimensions[r].height = 30
    r += 1
r += 1
warn2 = ws2.cell(row=r, column=1, value="Amoniakinhibitie-kanttekening: het 'geschat biogaspotentieel' voor pluimveemest is een theoretisch "
    "volume x opbrengstfactor-cijfer. Puur mono-vergisten van 100% pluimveemest kan in de praktijk leiden tot amoniakinhibitie van het "
    "vergistingsproces (hoog N-gehalte/laag C:N); in de praktijk wordt pluimveemest doorgaans verdund of meevergist met rundveemest. Dit "
    "cijfer is dus GEEN procesontwerp voor een mono-vergister op 100% pluimveemest - vandaar de kolom 'amoniakinhibitie_waarschuwing' in de "
    "CSV (JA wanneer een bedrijf >50% van zijn mestvolume uit pluimveemest heeft).")
warn2.font = WARN_FONT
warn2.fill = WARN_FILL
warn2.alignment = WRAP
ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws2.row_dimensions[r].height = 65
r += 2
loc_note = ws2.cell(row=r, column=1, value="Locatiemethode (v1.1, 29-09-2026): elk fictief bedrijf krijgt een VOLLEDIG willekeurige lat/lon "
    "binnen een schematische (niet-kadastrale) Nijkerk-omtrek, met het bebouwde centrum uitgesloten en minimaal 150 m tussen elk paar punten. "
    "Er wordt GEEN enkele echte geodataset (BRP-perceel, BAG-gebouw, adres) meer gebruikt om een locatie te bepalen - een eerdere versie deed "
    "dat wel (echte BRP-perceelcentroid + 80-250 m verschuiving) maar dat werd na review als onvoldoende privacy-veilig voor een "
    "gemeentepresentatie beoordeeld en losgelaten. "
    "afstand_arkemheen_m is de afstand (haversine) tot een INDICATIEF, zelfgekozen referentiepunt bij Arkemheen "
    "(52,265 N / 5,405 O) - GEEN officiele Natura 2000-gebiedsgrens-polygon.")
loc_note.font = NOTE_FONT
loc_note.alignment = WRAP
ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws2.row_dimensions[r].height = 65
r += 2
note2 = ws2.cell(row=r, column=1, value="CBS-bron: StatLine dataset 80781ned 'Landbouw; gewassen, dieren en grondgebruik naar gemeente', "
    "regio GM0267 (Nijkerk), periode 2023JJ00, via OData API, geraadpleegd 29-09-2026. "
    "RVO-bron: Tabel 4 'Diergebonden normen 2024' en Tabel 6 'Stikstof en fosfaat per melkkoe 2024', rvo.nl/onderwerpen/mest/tabellen.")
note2.font = NOTE_FONT
note2.alignment = WRAP
ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws2.row_dimensions[r].height = 45

# --- Sheet 3: sample of the synthetic data ---
ws3 = wb.create_sheet("Voorbeeld synthetische data")
with open("../data/nijkerk_synthetic_farms_2023.csv") as f:
    reader = csv.reader(f)
    for ri, row in enumerate(reader, start=1):
        for ci, val in enumerate(row, start=1):
            c = ws3.cell(row=ri, column=ci, value=val)
            if ri == 1:
                c.font = HEADER_FONT
                c.fill = HEADER_FILL
            else:
                c.font = BODY
        if ri > 20:
            break
for col_letter in "ABCDEFGHIJKLMNOPQRSTU":
    ws3.column_dimensions[col_letter].width = 13
ws3.cell(row=23, column=1, value="(volledige 75-rijen dataset: zie data/nijkerk_synthetic_farms_2023.csv)").font = NOTE_FONT

wb.save("../docs/Nijkerk_Synthetische_Demodata_Documentatie.xlsx")
print("saved doc workbook")
