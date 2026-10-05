"""Read the Emissieregistratie export and keep the NH3 and NOx emissions for Nijkerk.

The Emissieregistratie has no download API, so the export is made by hand:
    1. Open https://www.emissieregistratie.nl/data/data-export and choose "Export 1990-2024".
    2. Compartiment: Lucht
    3. Stof: Ammoniak and Stikstofoxiden (als NO2), each once, under "Alle stoffen"
    4. Gebiedsindeling: Gemeente
    5. Bronniveau: Sector
    6. Jaren: all years
    7. Download the XLSX and put it in data/nitrogen/raw/

Usage (from the repo root):
    pip install openpyxl
    python scripts/nitrogen/load_emissions.py

Writes data/nitrogen/processed/emissions_by_sector.json

The figures are copied as published. Emissions (kg of a substance released here)
are a different quantity from deposition (mol N arriving here); nothing is
converted from one to the other.
"""
from __future__ import annotations

import json
import sys
import warnings
from datetime import date

from openpyxl import load_workbook

import config


def find_export():
    files = sorted(config.RAW_DIR.glob(config.ER_EXPORT_PATTERN), key=lambda p: p.stat().st_mtime)
    if not files:
        sys.exit(f"No {config.ER_EXPORT_PATTERN} in data/nitrogen/raw/. Make the export by hand "
                 f"(steps are at the top of this script): {config.ER_EXPORT_PAGE}")
    return files[-1]


def read_sheet(workbook, name: str) -> list[tuple]:
    sheet = workbook[name]
    sheet.reset_dimensions()  # the export declares a wrong size; without this only one cell is read
    return [row for row in sheet.iter_rows(values_only=True) if any(cell is not None for cell in row)]


def read_metadata(rows: list[tuple]) -> dict:
    values = {str(row[0]): row[1] for row in rows if len(row) > 1 and row[1] is not None}
    cells = [str(cell).strip() for row in rows for cell in row if cell is not None]
    terms = next((text for text in cells if "bronvermelding" in text), None)
    return {"source_url": values.get("Bron"), "exported_at": values.get("Tijdstip"), "terms_nl": terms}


def main() -> None:
    path = find_export()
    print(f"Reading {path.name} ...")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl warns that the file has no default style
        workbook = load_workbook(path, read_only=True, data_only=True)
    metadata = read_metadata(read_sheet(workbook, "Metadata"))
    rows = read_sheet(workbook, "Emissies")
    header = {name: index for index, name in enumerate(rows[0])}
    needed = ["Stofcode", "Code_gebied", "Gebied", "Sector", "Jaar", "Emissie", "Eenheid", "Dataset", "Compartiment", "Brontype"]
    missing = [name for name in needed if name not in header]
    if missing:
        sys.exit(f"The export misses column(s): {', '.join(missing)}")

    def cell(row, name):
        return row[header[name]]

    area_rows = [row for row in rows[1:] if str(cell(row, "Code_gebied")) == config.ER_AREA_CODE]
    if not area_rows:
        sys.exit(f"No rows for area code {config.ER_AREA_CODE}. Was 'Gemeente' chosen as Gebiedsindeling?")
    if {cell(row, "Brontype") for row in area_rows} != {"Sector"}:
        sys.exit("The export is not at sector level. Choose 'Sector' as Bronniveau.")
    units = {cell(row, "Eenheid") for row in area_rows}
    if units != {"kg"}:
        sys.exit(f"Unexpected unit(s) in the export: {units}")

    years = sorted({int(cell(row, "Jaar")) for row in rows[1:]})
    national: dict[str, dict[int, float]] = {code: {} for code in config.ER_SUBSTANCES}
    for row in rows[1:]:
        code = str(cell(row, "Stofcode"))
        if code in national and str(cell(row, "Code_gebied")) not in config.ER_NON_MUNICIPAL_CODES:
            year = int(cell(row, "Jaar"))
            national[code][year] = national[code].get(year, 0.0) + float(cell(row, "Emissie"))
    substances = {}
    for code, meta in config.ER_SUBSTANCES.items():
        own = [row for row in area_rows if str(cell(row, "Stofcode")) == code]
        if not own:
            sys.exit(f"Substance {meta['label_nl']} (code {code}) is not in the export.")
        by_sector: dict[str, dict[int, float]] = {}
        for row in own:
            by_sector.setdefault(cell(row, "Sector"), {})[int(cell(row, "Jaar"))] = float(cell(row, "Emissie"))
        latest = years[-1]
        dominant = max(by_sector, key=lambda sector: by_sector[sector].get(latest, 0.0))
        # A year without a figure for the dominant sector gives a total that means something else.
        incomplete = [year for year in years if year not in by_sector[dominant]]
        order = sorted(by_sector, key=lambda sector: -by_sector[sector].get(latest, 0.0))
        substances[meta["id"]] = {
            "label_nl": meta["label_nl"], "formula": meta["formula"], "unit_nl": meta["unit_nl"], "unit": "kg",
            "dominant_sector": dominant,
            "years_without_dominant_sector": incomplete,
            "sectors": {sector: [(round(by_sector[sector][y], 1) if y in by_sector[sector] else None) for y in years]
                        for sector in order},
            "total": [None if y in incomplete else round(sum(v.get(y, 0.0) for v in by_sector.values()), 1) for y in years],
            # Context only: the same substance summed over all municipalities in the export.
            "all_municipalities_total": [None if y in incomplete else round(national[code].get(y, 0.0), 1) for y in years],
        }

    output = {
        "source": {
            "publisher": "Emissieregistratie (RIVM, CBS, PBL, WUR, Deltares)",
            "dataset": sorted({str(cell(row, "Dataset")) for row in area_rows})[0],
            "url": config.ER_EXPORT_PAGE,
            "export_file": path.name,
            "exported_at": metadata["exported_at"],
            "selection": "Compartiment Lucht, Gebiedsindeling Gemeente, Bronniveau Sector",
            "terms_nl": metadata["terms_nl"],
            "processed_on": date.today().isoformat(),
        },
        "notes": [
            "Figures are copied from the export; only the unit is kept as kg and rounded to 0.1 kg.",
            "Emissions are assigned to the place where they are released. They are not deposition.",
            "NOx is expressed as NO2 by the Emissieregistratie.",
            "A missing value means the export has no row for that sector and year; it is not a zero.",
            "The total is left empty for years in which the dominant sector has no figure.",
            "all_municipalities_total sums every municipality in the export (the area 'Noordzee' is left out). It is context, not an official national total.",
            "Each yearly release of the Emissieregistratie recalculates earlier years. Do not mix releases.",
        ],
        "area": {"code": config.ER_AREA_CODE, "name": str(cell(area_rows[0], "Gebied"))},
        "years": years,
        "substances": substances,
    }
    config.PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    target = config.PROCESSED_DIR / "emissions_by_sector.json"
    target.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {target.relative_to(config.REPO_ROOT)} ({output['source']['dataset']}, {years[0]}-{years[-1]})")

    for key, sub in substances.items():
        print(f"\n{output['area']['name']}, {sub['label_nl']} ({sub['formula']}), tonnes per year")
        shown = [y for y in years if y in (years[0], 2000, 2015, years[-1])]
        print(f"  {'Sector':44}" + "".join(f"{y:>9}" for y in shown))
        for sector, values in sub["sectors"].items():
            print(f"  {sector[:44]:44}" + "".join(f"{'-':>9}" if values[years.index(y)] is None else f"{values[years.index(y)] / 1000:9.1f}" for y in shown))
        print(f"  {'Total':44}" + "".join(f"{'-':>9}" if sub["total"][years.index(y)] is None else f"{sub['total'][years.index(y)] / 1000:9.1f}" for y in shown))
        if sub["years_without_dominant_sector"]:
            print(f"  No figure for {sub['dominant_sector']} in: {sub['years_without_dominant_sector']} (total left empty)")


if __name__ == "__main__":
    main()
